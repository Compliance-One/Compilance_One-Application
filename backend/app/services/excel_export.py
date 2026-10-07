"""
Builds the Excel workbook from already-fetched P&L / balance sheet dicts
(see reports.py's /export/excel route). Kept synchronous on purpose —
openpyxl itself is sync, and this way it never needs its own DB session.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

BOLD = Font(bold=True)


def build_financial_statements_xlsx(pl: dict, bs: dict) -> bytes:
    wb = Workbook()

    ws_pl = wb.active
    ws_pl.title = "Profit and Loss"
    ws_pl.append(["Profit & Loss", f"{pl['period_from']} to {pl['period_to']}"])
    ws_pl["A1"].font = BOLD
    ws_pl.append([])
    ws_pl.append(["Income by category"])
    ws_pl["A4"].font = BOLD
    for category, amount in pl["income_by_category"].items():
        ws_pl.append([category, amount])
    ws_pl.append(["Total income", pl["total_income"]])
    ws_pl.append([])
    ws_pl.append(["Expense by category"])
    for category, amount in pl["expense_by_category"].items():
        ws_pl.append([category, amount])
    ws_pl.append(["Total expense", pl["total_expense"]])
    ws_pl.append([])
    ws_pl.append(["Net profit", pl["net_profit"]])
    ws_pl["A" + str(ws_pl.max_row)].font = BOLD

    ws_bs = wb.create_sheet("Balance Sheet")
    ws_bs.append(["Balance Sheet"])
    ws_bs["A1"].font = BOLD
    ws_bs.append([])
    ws_bs.append(["Cash & bank", bs["cash_and_bank"], "(not yet tracked in the schema)"])
    ws_bs.append(["Stock value", bs["stock_value"]])
    ws_bs.append(["Receivables", bs["receivables"]])
    ws_bs.append(["Payables", bs["payables"]])
    ws_bs.append(["Owner's capital", bs["owners_capital"]])

    for ws in (ws_pl, ws_bs):
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 18

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
